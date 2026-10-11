"""Oppmøte-eksport fra Spond: lagre fila og lag et kontaktfritt sammendrag.

Sponds eksport (``sponds/{id}/export``) har to ark. «For print» har én blokk
per svarstatus; «For import» har én rad per medlem med kolonnene Status, Navn,
E-post, Mobilnummer, Foresatt N's navn/e-post/mobilnummer og Svartidspunkt.
Tittelradene øverst er overskrift, tidspunkt og sted.

Fila på disk beholder alt Spond sender (også kontaktdata). Sammendraget som
går tilbake til Claude har bare status, navn og svartidspunkt.
"""

from __future__ import annotations

import io
import os
import re
import warnings
from collections import Counter
from datetime import datetime
from pathlib import Path

CONTACT_RE = re.compile(r"e-?post|mail|telefon|phone|mobil|adresse|address|fødsel|birth|foresatt", re.I)
KEEP = ("Status", "Navn", "Svartidspunkt")


def export_dir() -> Path:
    """``SPOND_MCP_EXPORT_DIR``, ellers ``~/Spond-eksport`` (ikke Dokumenter/Skrivebord,
    som macOS ber om tillatelse til for apper som Claude Desktop)."""
    env = os.environ.get("SPOND_MCP_EXPORT_DIR")
    return Path(env).expanduser() if env else Path.home() / "Spond-eksport"


def _slug(text: str) -> str:
    s = re.sub(r"[^\wæøåÆØÅ]+", "-", (text or "arrangement").lower()).strip("-")
    return s[:40] or "arrangement"


def save(data: bytes, directory: Path, date_str: str, heading: str) -> Path:
    """Skriv fila uten å overskrive (``_2``, ``_3`` …)."""
    directory.mkdir(parents=True, exist_ok=True)
    base = f"spond-oppmote-{date_str}-{_slug(heading)}"
    path = directory / f"{base}.xlsx"
    n = 2
    while path.exists():
        path = directory / f"{base}_{n}.xlsx"
        n += 1
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)  # kontaktdata: bare eier
    with os.fdopen(fd, "wb") as fh:
        fh.write(data)
    return path


def summarize(data: bytes) -> dict:
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # «Workbook contains no default style»
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = list(wb.sheetnames)
    try:
        ws = wb["For import"] if "For import" in sheets else wb.worksheets[-1]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    title = [str(r[0]) for r in rows[:3] if r and r[0]]
    header_idx = next(
        (i for i, r in enumerate(rows) if r and "Status" in r and "Navn" in r), None
    )
    if header_idx is None:
        return {"sheets": sheets, "title": title, "rows": len(rows),
                "warning": "Fant ikke kolonnene Status/Navn; se fila."}

    header = [str(h) if h is not None else "" for h in rows[header_idx]]
    idx = {name: header.index(name) for name in KEEP if name in header}
    dropped = sorted({h for h in header if h and h not in KEEP and CONTACT_RE.search(h)})

    by_status: dict[str, list[str]] = {}
    answered_at = []
    for r in rows[header_idx + 1:]:
        if not r or r[idx["Navn"]] in (None, ""):
            continue
        status = str(r[idx["Status"]] or "ukjent")
        by_status.setdefault(status, []).append(str(r[idx["Navn"]]))
        if "Svartidspunkt" in idx and isinstance(r[idx["Svartidspunkt"]], datetime):
            answered_at.append(r[idx["Svartidspunkt"]])

    return {
        "title": title,
        "columns_kept": [k for k in KEEP if k in idx],
        "columns_dropped": dropped,
        "rows": sum(len(v) for v in by_status.values()),
        "by_status": dict(Counter({k: len(v) for k, v in by_status.items()})),
        "names_by_status": {k: sorted(v) for k, v in by_status.items()},
        "first_answer": min(answered_at).strftime("%Y-%m-%d %H:%M") if answered_at else None,
        "last_answer": max(answered_at).strftime("%Y-%m-%d %H:%M") if answered_at else None,
    }
